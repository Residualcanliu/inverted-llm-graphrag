"""读取原始文件。

每个函数负责一类文件，返回统一的结构化字典。规范化交给 normalize.py，
这里只做「读出来 + 按字段名整理」。
"""

from __future__ import annotations

import glob
import os
import re

import docx
import pandas as pd

from app.ingest.normalize import (
    clean_text,
    expand_ids,
    parse_date,
    parse_generic_models,
    split_name_spec,
    split_multi,
)

RAW_SUBDIR = "设备表备份_2026-09-23"


def _raw_dir(raw_dir: str) -> str:
    """原始文件可能直接放在 raw/ 下，也可能在带日期的子目录里。"""
    sub = os.path.join(raw_dir, RAW_SUBDIR)
    return sub if os.path.isdir(sub) else raw_dir


# ---------------- 设备台账 ----------------

def read_equipment(raw_dir: str) -> list[dict]:
    df = pd.read_excel(os.path.join(_raw_dir(raw_dir), "设备1.xlsx"))
    out = []
    for _, r in df.iterrows():
        eid = clean_text(r.get("设备编号"))
        if not eid:
            continue
        out.append({
            "id": eid,
            "model": eid.split("-")[0],
            "shop": clean_text(r.get("分配单位")),
            "function": clean_text(r.get("功能")),
            "tool_groups": [x for x in split_multi(r.get("设备适配工具组"))
                            if x != "通用工具"],
            "uses_common_tools": "通用工具" in clean_text(r.get("设备适配工具组")),
            "warehouse": clean_text(r.get("备件仓库")),
            "inspected": clean_text(r.get("是否检修")) == "已检修",
            "priority": clean_text(r.get("重点标记")),
        })
    return out


# ---------------- 备件与工具（合并成 SparePart） ----------------

def read_spare_parts(raw_dir: str) -> list[dict]:
    """备件和工具读成同一类。

    两者在源数据里结构几乎一样，合并后用 category 字段区分。
    差异是备件有存放位置和库存，工具有数量。
    """
    root = _raw_dir(raw_dir)
    out: list[dict] = []

    for f in sorted(glob.glob(os.path.join(root, "备件仓库", "*.xlsx"))):
        for _, r in pd.read_excel(f).iterrows():
            rid = clean_text(r.get("备件编号"))
            if not rid:
                continue
            name, spec = split_name_spec(r.get("备件名称"))
            out.append({
                "id": rid,
                "name": name,
                "spec": spec,
                "category": "备件",
                "warehouse": os.path.basename(f).replace("设备备件仓库", "").replace(".xlsx", ""),
                "location": clean_text(r.get("备件存放位置")),
                "stock": _num(r.get("库存数量")),
                "safety_stock": _num(r.get("安全库存")),
                "fits_models": parse_generic_models(r.get("通用设备")),
                "note": clean_text(r.get("标注")),
            })

    for f in sorted(glob.glob(os.path.join(root, "工具组表", "*.xlsx"))):
        group = os.path.basename(f).replace(".xlsx", "")
        for _, r in pd.read_excel(f).iterrows():
            rid = clean_text(r.get("工具编号"))
            if not rid:
                continue
            name, spec = split_name_spec(r.get("工具名称"))
            note = clean_text(r.get("备注"))
            out.append({
                "id": rid,
                "name": name,
                "spec": spec or clean_text(r.get("规格型号")),
                "category": "工具",
                "group": group,
                "location": "",
                "stock": _num(r.get("数量")),
                "safety_stock": None,
                "fits_models": parse_generic_models(r.get("适配设备")),
                "note": note,
                "hazard": "高危" in note,
                "consumable": "一次性消耗" in note,
            })
    return out


def split_causes(text: str) -> list[str]:
    """把成因列拆成独立的成因。

    源数据的「快速判断」列常用「A或B」的写法表达多个可能成因，
    例如「丝杠间隙过大或伺服报警」。整条留着一个名字，会带来两个问题：
    它跟真正的故障名判成上下位关系；藏在里面的「伺服报警」暴露不出来。

    拆开之后，某些成因正好就是检修表里的故障名，TRIGGERS 边才建得起来。
    """
    s = clean_text(text)
    if not s:
        return []
    parts = re.split(r"[或，,;；]|以及|并且", s)
    return [p.strip() for p in parts if p.strip()]


def _num(v):
    if v is None:
        return None
    try:
        f = float(v)
        return None if f != f else int(f)      # NaN 判断
    except (TypeError, ValueError):
        return None


# ---------------- 检修记录 ----------------

def read_repairs(raw_dir: str) -> list[dict]:
    df = pd.read_excel(os.path.join(_raw_dir(raw_dir), "设备检修表.xlsx"))
    out = []
    for _, r in df.iterrows():
        dev = clean_text(r.get("设备编号"))
        if not dev:
            continue
        d = parse_date(r.get("检修日期"))
        out.append({
            "device": dev,
            "date": d.isoformat() if d else "",
            "failure": clean_text(r.get("故障描述")),
            "parts": split_multi(r.get("更换配件")),
            "tools": split_multi(r.get("使用工具")),
            "done": clean_text(r.get("检修是否完成")) == "已完成",
        })
    return out


# ---------------- 工艺流程图 ----------------

def read_processes(raw_dir: str) -> tuple[list[dict], list[str]]:
    """返回 (流程列表, 自检信息)。

    第二张表是流程 × 工序明细，每格标了台数，用来做展开后的数量校验。
    """
    d = docx.Document(os.path.join(_raw_dir(raw_dir), "工艺流程图.docx"))
    if len(d.tables) < 2:
        return [], ["工艺流程图表格数不足，跳过"]

    # 表1：流程总览
    overview = {}
    for row in [[c.text.strip() for c in r.cells] for r in d.tables[0].rows][1:]:
        if len(row) >= 5 and row[0]:
            overview[row[0]] = {"name": row[1], "chain": row[2],
                                "part_type": row[4] if len(row) > 4 else ""}

    # 表2：流程 × 工序明细
    from app.ingest.normalize import check_count, extract_count
    flows: dict[str, dict] = {}
    checks: list[str] = []
    for row in [[c.text.strip() for c in r.cells] for r in d.tables[1].rows][1:]:
        if len(row) < 5 or not row[0]:
            continue
        fid, fname, order, step, devices = row[0], row[1], row[2], row[3], row[4]
        note = row[5] if len(row) > 5 else ""
        devs = expand_ids(devices)
        label = f"{fid} {step}"
        try:
            checks.append(check_count(extract_count(devices), devs, label))
        except Exception as e:                       # noqa: BLE001
            checks.append(f"[失败] {e}")

        f = flows.setdefault(fid, {
            "id": fid, "name": fname,
            "part_type": overview.get(fid, {}).get("part_type", ""),
            "steps": [],
        })
        n = int(order.split("/")[0]) if "/" in order else len(f["steps"]) + 1
        f["steps"].append({
            "order": n,
            "operation": step.split("（")[0].strip(),
            "batch": step[step.find("（") + 1:step.find("）")] if "（" in step else "",
            "devices": devs,
            "note": note,
        })

    for f in flows.values():
        f["steps"].sort(key=lambda s: s["order"])
    return list(flows.values()), checks


# ---------------- 故障处理手册 ----------------

def read_manuals(raw_dir: str) -> list[dict]:
    out = []
    for f in sorted(glob.glob(os.path.join(_raw_dir(raw_dir), "故障紧急处理文件", "*.docx"))):
        d = docx.Document(f)
        m = re.search(r"\((\w+)\)", os.path.basename(f))
        model = m.group(1) if m else ""
        title = next((p.text.strip() for p in d.paragraphs if p.text.strip()), "")

        failures = []
        for tb in d.tables:
            for row in [[c.text.strip() for c in r.cells] for r in tb.rows][1:]:
                if len(row) >= 3 and row[0]:
                    failures.append({
                        "symptom": row[0],
                        "cause": row[1],
                        "causes": split_causes(row[1]),
                        "action": row[2],
                    })

        # 安全注意事项：在「安全注意事项」标题之后，逐条取
        safety, capturing = [], False
        for p in d.paragraphs:
            t = p.text.strip()
            if not t:
                continue
            if "安全注意" in t:
                capturing = True
                continue
            if capturing:
                if t.startswith(("一、", "二、", "三、", "四、", "五、", "六、")):
                    break
                safety.append(re.sub(r"^[·•\-\*]\s*", "", t))

        out.append({
            "model": model,
            "title": title,
            "failures": failures,
            "safety": safety,
        })
    return out
