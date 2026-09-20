"""把 DESIGN.md 转成排版好的 Word 文档。

改完 DESIGN.md 后重跑这个脚本即可，不用手工同步两份文档。

用法：python scripts/md2docx.py [源文件] [输出文件]
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:                                     # noqa: BLE001
    pass

import docx                                                     # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT                  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH                   # noqa: E402
from docx.oxml import OxmlElement                               # noqa: E402
from docx.oxml.ns import qn                                     # noqa: E402
from docx.shared import Cm, Pt, RGBColor                        # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "DESIGN.md"
OUT = (Path(sys.argv[2]) if len(sys.argv) > 2 else
       ROOT / "docs" / "技术设计文档.docx")

BODY, MONO = "微软雅黑", "Consolas"
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
INLINE_CODE = RGBColor(0xB0, 0x30, 0x60)

doc = docx.Document()
_n = doc.styles["Normal"]
_n.font.name = BODY
_n.font.size = Pt(10.5)
_n.element.rPr.rFonts.set(qn("w:eastAsia"), BODY)
_n.paragraph_format.space_after = Pt(6)
_n.paragraph_format.line_spacing = 1.5
for _s in doc.sections:
    _s.top_margin = _s.bottom_margin = Cm(2.2)
    _s.left_margin = _s.right_margin = Cm(2.2)


def sf(run, name=BODY, size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    return run


def shade(p, fill):
    e = OxmlElement("w:shd")
    e.set(qn("w:val"), "clear")
    e.set(qn("w:fill"), fill)
    p._p.get_or_add_pPr().append(e)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def add_inline(p, text, size=10.5, bold=False):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            sf(p.add_run(part[2:-2]), BODY, size, True)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            sf(p.add_run(part[1:-1]), MONO, size - 0.5, False, INLINE_CODE)
        elif part.startswith("[") and "](" in part:
            sf(p.add_run(part[1:part.index("](")]), BODY, size, bold)
        else:
            sf(p.add_run(part), BODY, size, bold)
    # 兜底：落单的 ** 直接清掉，避免漏进 Word
    for r in p.runs:
        if "**" in r.text:
            r.text = r.text.replace("**", "")


def heading(text, level):
    sizes = {1: 19, 2: 14.5, 3: 12, 4: 11}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4 if level == 1 else 15)
    p.paragraph_format.space_after = Pt(6)
    if level <= 2:
        p.paragraph_format.keep_with_next = True
    sf(p.add_run(text), BODY, sizes.get(level, 11), True,
       ACCENT if level <= 2 else RGBColor(0x1A, 0x1A, 0x1A))


def make_table(rows):
    hdr = [c.strip() for c in rows[0].strip("|").split("|")]
    body = [[c.strip() for c in r.strip("|").split("|")] for r in rows[2:]]
    t = doc.add_table(rows=1, cols=len(hdr))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(hdr):
        c = t.rows[0].cells[i]
        c.text = ""
        add_inline(c.paragraphs[0], h, 9, True)
        shade(c.paragraphs[0], "DCE6F1")
    for row in body:
        cells = t.add_row().cells
        for i in range(len(hdr)):
            cells[i].text = ""
            add_inline(cells[i].paragraphs[0], row[i] if i < len(row) else "", 9)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def code_block(lines):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.1
    for i, ln in enumerate(lines):
        if i:
            p.add_run().add_break()
        sf(p.add_run(ln), MONO, 8.5, False, RGBColor(0x1A, 0x1A, 0x1A))
    shade(p, "F4F6F8")


def emit_quote(lines):
    """一个引用段落（已被空 > 行切开）。可能是列表。"""
    stripped = [l.strip() for l in lines if l.strip()]
    if stripped and all(re.match(r"^[-*]\s+", l) for l in stripped):
        for b in stripped:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Cm(1.1)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.35
            add_inline(p, re.sub(r"^[-*]\s+", "", b), 10)
        return
    txt = " ".join(stripped)
    if not txt:
        return
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.7)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.35
    add_inline(p, txt, 10)


def main() -> int:
    raw = SRC.read_text(encoding="utf-8").split("\n")
    i, first_h1 = 0, False

    while i < len(raw):
        s = raw[i].strip()

        if s.startswith("```"):
            i += 1
            buf = []
            while i < len(raw) and not raw[i].strip().startswith("```"):
                buf.append(raw[i])
                i += 1
            i += 1
            code_block(buf)
            continue

        if s.startswith("|") and i + 1 < len(raw) and \
                re.match(r"^\|[\s:\-|]+\|$", raw[i + 1].strip()):
            rows = []
            while i < len(raw) and raw[i].strip().startswith("|"):
                rows.append(raw[i].strip())
                i += 1
            make_table(rows)
            continue

        if s == "---":
            i += 1
            continue

        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            lvl, txt = len(m.group(1)), m.group(2).strip()
            if lvl == 1 and not first_h1:
                first_h1 = True
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(14)
                sf(p.add_run(txt), BODY, 22, True, ACCENT)
            else:
                heading(txt, lvl)
            i += 1
            continue

        if s.startswith(">"):
            blocks, cur = [], []
            while i < len(raw) and raw[i].strip().startswith(">"):
                content = raw[i].strip().lstrip(">").strip()
                if content:
                    cur.append(content)
                elif cur:
                    blocks.append(cur)
                    cur = []
                i += 1
            if cur:
                blocks.append(cur)
            for b in blocks:
                emit_quote(b)
            continue

        m = re.match(r"^[-*]\s+(.*)$", s)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.4
            add_inline(p, m.group(1), 10.5)
            i += 1
            continue

        m = re.match(r"^\d+\.\s+(.*)$", s)
        if m:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.4
            add_inline(p, m.group(1), 10.5)
            i += 1
            continue

        if not s:
            i += 1
            continue

        p = doc.add_paragraph()
        add_inline(p, s, 10.5)
        i += 1

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)

    # 自检：不能有 markdown 标记漏进 Word
    leaked = [p.text for p in doc.paragraphs if "**" in p.text or "](" in p.text]
    print(f"已生成 {OUT.name}")
    print(f"  段落 {len(doc.paragraphs)}  表格 {len(doc.tables)}")
    if leaked:
        print(f"  [警告] {len(leaked)} 段残留 markdown 标记：")
        for t in leaked[:3]:
            print(f"    {t[:70]}")
        return 1
    print("  无残留标记")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
